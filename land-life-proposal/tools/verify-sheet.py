#!/usr/bin/env python3
"""スプレッドシート版の数式が lib/engine.js と同じ結果になるか確かめる（LibreOffice で再計算して比較）。

  python3 tools/build-sheet.py && python3 tools/verify-sheet.py
"""
import json
import os
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

JS = r"""
const LL=require('./lib/engine'),E=require('./editions');
const out={};const C={};
for(const [id,ed] of Object.entries(E)){
  const x=LL.compute(ed.sample);
  out[id]={label:ed.label,comfort:ed.comfortRatio,savings:x.c.savings,tax:x.tax,rent:x.rent.total,afford:x.afford,props:x.props.map(p=>({
    loan:p.loan.amount,pmt:p.pmt,cashOut:p.cashOut,cashLeft:p.cashLeft,own:p.own.total,
    ratio:(p.pmt*12+(x.c.existingMonthly||0)*12)/x.c.gross,screen:p.screen,
    ok:p.credit.judge.ok,limit:p.credit.judge.limit,years:p.credit.judge.years,
    rows:p.credit.rows.slice(0,13).map(r=>({bal:r.balance,max:r.max,it:r.fromIT,rt:r.fromRT,act:r.actual})),
    total:p.credit.total,limitScreen:p.screen.limit,totalMax:p.credit.totalMax,first10:p.credit.first10,ten:p.ten.with,
    risk:p.risk.map(r=>({th:r.takeHomeM,gap:r.gap})),
    sale:p.sale.map(s=>s.byScen.map(b=>b.net))}))};
}
console.log(JSON.stringify(out));
"""


def recalc(path, tmp):
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, path],
                   check=True, capture_output=True)
    return load_workbook(os.path.join(tmp, os.path.basename(path)), data_only=True)


def find(ws, label, col=1, start=1):
    for r in range(start, ws.max_row + 1):
        if ws.cell(r, col).value == label:
            return r
    raise KeyError(f"{ws.title}: {label}")


def main():
    eng = json.loads(subprocess.check_output(["node", "-e", JS], cwd=ROOT))
    fails, checks = [], 0

    def eq(name, got, want, tol=1.0):
        nonlocal checks
        checks += 1
        if isinstance(want, bool) or isinstance(got, str):
            ok = got == want
        else:
            ok = got is not None and abs(got - want) <= tol
        if not ok:
            fails.append(f"{name}: sheet={got!r} engine={want!r}")

    with tempfile.TemporaryDirectory() as tmp:
        for id_, e in eng.items():
            path = os.path.join(ROOT, "dist", id_, f"資金計画シミュレーター_{e['label']}.xlsx")
            wb = recalc(path, tmp)
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        if isinstance(c.value, str) and c.value.startswith("#") or c.value in ("Err:502", "Err:504", "#NAME?"):
                            fails.append(f"{id_} {ws.title}!{c.coordinate} = {c.value}")
            raw = load_workbook(path)
            # わかりやすさの受け入れ基準（docs/05_スプレッドシート版の設計.md §5）
            first = raw.worksheets[0]
            checks += 1
            if first.title != "① かんたん試算":
                fails.append(f"{id_} 1枚目が「① かんたん試算」ではない: {first.title}")
            yellow = sum(1 for row in first.iter_rows() for cl in row
                         if cl.fill is not None and cl.fill.fgColor is not None and cl.fill.fgColor.rgb in ("00FFF2CC", "FFFFF2CC"))
            checks += 1
            if not 0 < yellow <= 15:
                fails.append(f"{id_} かんたん試算の入力欄が {yellow} か所（1〜15か所にする）")
            # 第2版（§6）：入力欄はすべてプルダウン・完成度100%・未入力なら結果を出さない
            ycells = [cl for row in first.iter_rows() for cl in row
                      if cl.fill is not None and cl.fill.fgColor is not None and cl.fill.fgColor.rgb in ("00FFF2CC", "FFFFF2CC")]
            listed = [str(dv.sqref) for dv in first.data_validations.dataValidation if dv.type == "list"]
            for cl in ycells:
                checks += 1
                if not any(cl.coordinate in rng.split() for rng in listed):
                    fails.append(f"{id_} {cl.coordinate} の入力欄がプルダウンではない")
            k0 = wb["① かんたん試算"]
            prog = k0.cell(find(k0, "入力の完成度"), 2).value or ""
            checks += 1
            if "100%" not in prog:
                fails.append(f"{id_} サンプルの完成度が100%でない: {prog!r}")
            blank = os.path.join(tmp, "blank", os.path.basename(path))
            os.makedirs(os.path.dirname(blank), exist_ok=True)
            bw = load_workbook(path)
            bw.worksheets[0][ycells[0].coordinate].value = None
            bw.save(blank)
            bwb = recalc(blank, os.path.join(tmp, "blank", "out"))
            b0 = bwb["① かんたん試算"]
            checks += 3
            if "あと1項目" not in str(b0.cell(find(b0, "入力の完成度"), 4).value):
                fails.append(f"{id_} 1つ空にしても「あと1項目」と出ない: {b0.cell(find(b0, '入力の完成度'), 4).value!r}")
            if b0.cell(find(b0, "毎月の住まいの支払い"), 3).value not in (None, ""):
                fails.append(f"{id_} 未入力なのに結果が出ている")
            if "すべて入力すると" not in str(b0.cell(find(b0, "物件A"), 2).value):
                fails.append(f"{id_} 未入力なのにまとめが出ている")
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cl in row:
                        if isinstance(cl.value, bool):
                            fails.append(f"{id_} {ws.title}!{cl.coordinate} に TRUE/FALSE が表示される")
            m, y, z, v, t = wb["③ 詳しい試算"], wb["④ 住宅ローン控除"], wb["⑤ 収入が減ったら"], wb["⑥ 売却の目安"], wb["参考：税額の概算"]
            k_ = wb["① かんたん試算"]
            kv = lambda label, col: k_.cell(find(k_, label), col).value
            eq(f"{id_} かんたん 賃貸の住居費", kv("毎月の住まいの支払い", 2), e["rent"], 0.01)
            eq(f"{id_} かんたん 手元の預貯金（賃貸）", kv("手元に残る預貯金", 2), e["savings"] / 1e4, 1e-6)
            eq(f"{id_} かんたん 10年の賃貸", kv("10年間の住まいの費用", 2), e["rent"] * 120 / 1e4, 1e-4)
            for k, p in enumerate(e["props"]):
                n, col = f"{id_} かんたん 物件{'AB'[k]}", 3 + k
                eq(n + " 住居費", kv("毎月の住まいの支払い", col), p["own"], 0.01)
                eq(n + " 返済", kv("うちローンの返済", col), p["pmt"], 0.01)
                eq(n + " 初期費用", kv("はじめに必要な現金", col), p["cashOut"] / 1e4, 1e-4)
                eq(n + " 手元", kv("手元に残る預貯金", col), p["cashLeft"] / 1e4, 1e-4)
                eq(n + " 控除10年", kv("住宅ローン控除（最初の10年）", col), p["first10"] / 1e4, 1e-4)
                eq(n + " 10年総額", kv("10年間の住まいの費用", col), p["ten"] / 1e4, 1e-4)
                want = "◎" if p["ratio"] <= e["comfort"] else "△" if p["ratio"] <= p["limitScreen"] else "×"
                eq(n + " 返済の負担の判定", (kv("返済の負担", col) or "")[:1], want)
                eq(n + " 控除の判定", (kv("住宅ローン控除", col) or "")[:1], "◎" if p["ok"] else "×")
                eq(n + " 審査の判定", (kv("ローン審査の目安", col) or "")[:1], "◎" if p["screen"]["ratio"] <= p["limitScreen"] else "×")
            val = lambda ws, label, k, start=1: ws.cell(find(ws, label, 1, start), 2 + k).value
            eq(f"{id_} 所得税", val(t, "所得税（控除前）", 0), e["tax"]["incomeTax"])
            eq(f"{id_} 住民税", val(t, "住民税（所得割）", 0), e["tax"]["residentTax"])
            eq(f"{id_} 課税所得", val(t, "課税所得", 0), e["tax"]["taxable"])
            for k, p in enumerate(e["props"]):
                n = f"{id_} 物件{'AB'[k]}"
                eq(n + " 借入額", val(m, "借入額", k), p["loan"])
                eq(n + " 毎月の返済", val(m, "毎月の返済額", k), p["pmt"], 0.01)
                eq(n + " 初期費用", val(m, "初期費用（現金）", k), p["cashOut"])
                eq(n + " 住居費", val(m, "持ち家の住居費（月額）", k), p["own"], 0.01)
                eq(n + " 返済比率", val(m, "返済比率（年収比・既存の借入を含む）", k), p["ratio"], 1e-9)
                eq(n + " 審査比率", val(m, "審査の返済比率", k), p["screen"]["ratio"], 1e-9)
                eq(n + " 審査上限", val(m, "審査上の借入の上限（目安）", k), p["screen"]["maxLoan"])
                eq(n + " 控除対象", val(m, "控除の対象か", k), "はい" if p["ok"] else "いいえ")
                eq(n + " 限度額", val(m, "借入限度額", k), p["limit"])
                eq(n + " 期間", val(m, "控除期間（年）", k), p["years"])
                eq(n + " 控除合計", val(m, "控除見込額（期間合計・概算）", k), p["total"])
                eq(n + " 控除上限合計", val(m, "制度上の控除の上限（期間合計）", k), p["totalMax"])
                eq(n + " 10年控除", val(m, "最初の10年の控除見込額", k), p["first10"])
                eq(n + " 10年総額", val(m, "10年間の合計（控除あり）", k), p["ten"], 0.5)
                top = find(y, "年目", 1, find(y, f"物件{'AB'[k]}")) + 1
                for i, r in enumerate(p["rows"]):
                    eq(f"{n} {i + 1}年目 残高", y.cell(top + i, 3).value, r["bal"], 0.01)
                    eq(f"{n} {i + 1}年目 実際", y.cell(top + i, 9).value, r["act"])
                zs = find(z, f"物件{'AB'[k]}")
                for j, rk in enumerate(p["risk"]):
                    eq(f"{n} 収入減{j} 手取り", val(z, "手取り（月額・概算）", j, zs), rk["th"], 0.01)
                    eq(f"{n} 収入減{j} 過不足", val(z, "毎月の過不足", j, zs), rk["gap"], 0.01)
                for si, yrs in enumerate((5, 10, 20)):
                    vs = find(v, f"物件{'AB'[k]}・{yrs}年後に売る場合")
                    for j in range(3):
                        eq(f"{n} {yrs}年後 シナリオ{j}", val(v, "手元に残る額（マイナスは不足）", j, vs), p["sale"][si][j], 0.5)
            if e["afford"]:
                eq(f"{id_} 家賃並み価格", m.cell(find(m, "買える物件価格の目安"), 2).value, e["afford"]["price"])
            else:
                eq(f"{id_} 家賃並み価格", m.cell(find(m, "買える物件価格の目安"), 2).value, 0)
    print(f"checks: {checks}  fails: {len(fails)}")
    for f in fails[:40]:
        print("  ✗", f)
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
