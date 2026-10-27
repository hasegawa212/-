#!/usr/bin/env python3
"""資金計画シミュレーター（スプレッドシート版）を版ごとに作る。

  python3 tools/build-sheet.py            → dist/<版>/資金計画シミュレーター_<版名>.xlsx（全版）
  python3 tools/build-sheet.py martialarts

計算はすべてセルの数式で、lib/engine.js と同じ式を使う（住宅ローン・税額の概算・住宅ローン控除・審査の返済比率・
家賃並み予算・収入減少・売却の目安）。初期値は editions/*.js のサンプル。数値の一致は tools/verify-sheet.py で確認する。
Googleスプレッドシートにアップロードすると、そのまま数式つきで変換される。
"""
import json
import os
import re
import subprocess
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ── 見た目（全シート共通） ──
NAVY, GOLD = "1F2A44", "B08D57"
F_TITLE = Font(name="Arial", size=14, bold=True, color=NAVY)
F_NOTE = Font(name="Arial", size=9, color="595959")
F_HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_SEC = Font(name="Arial", size=11, bold=True, color=NAVY)
F_IN = Font(name="Arial", size=10, color="0000FF")
F_CALC = Font(name="Arial", size=10, color="000000")
F_LINK = Font(name="Arial", size=10, color="008000")
F_KEY = Font(name="Arial", size=11, bold=True, color="000000")
BG_HEAD = PatternFill("solid", fgColor=NAVY)
BG_IN = PatternFill("solid", fgColor="FFF2CC")
BG_KEY = PatternFill("solid", fgColor="F3EBDD")
LINE = Side(style="thin", color="BFBFBF")
BOX = Border(top=LINE, bottom=LINE, left=LINE, right=LINE)

YEN = '#,##0"円";[Red]-#,##0"円"'
MAN = '#,##0"万円";[Red]-#,##0"万円"'
MAN1 = '#,##0.0"万円";[Red]-#,##0.0"万円"'
SIGNED = '+#,##0"円";[Red]-#,##0"円";0"円"'
YEARS = '0"年"'
PCT = '0.0%'
PCT2 = '0.00%'
DATE = 'yyyy/mm/dd'
DATE0 = '[=0]"—";yyyy/mm/dd'  # 期限がない区分は0 → 「—」
NUM = '#,##0.0'

# タブ名（番号の順に見る。「参考：」は税率表など）
S_EASY, S_IN, S_CALC = "① かんたん試算", "② 詳しい条件", "③ 詳しい試算"
S_CR, S_RISK, S_SALE = "④ 住宅ローン控除", "⑤ 収入が減ったら", "⑥ 売却の目安"
S_TAX, S_TAXT, S_RULE = "参考：税額の概算", "参考：税の前提", "参考：控除の制度"
TAB = {S_EASY: GOLD, S_IN: "F1C232", S_CALC: NAVY, S_CR: NAVY, S_RISK: NAVY, S_SALE: NAVY, S_TAX: "A6A6A6", S_TAXT: "A6A6A6", S_RULE: "A6A6A6"}
BG_OK, BG_WARN, BG_NG = PatternFill("solid", fgColor="D9EAD3"), PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="F4CCCC")

KIND = {"new": "新築", "resale": "買取再販", "used": "中古"}
PERF = {"certified": "認定長期優良・低炭素", "zeh": "ZEH水準省エネ", "energy": "省エネ基準適合", "other": "その他"}


def editions():
    js = "const E=require('./editions');const L=require('./lib/rules');" \
         "console.log(JSON.stringify({E:Object.fromEntries(Object.entries(E).map(([k,v])=>[k,{id:v.id,label:v.label,brand:v.brand,sample:v.sample,comfortRatio:v.comfortRatio}])),R:L}))"
    out = subprocess.check_output(["node", "-e", js], cwd=ROOT)
    return json.loads(out)


def d(s):
    return date.fromisoformat(s) if s else None


class Sheet:
    """ラベル付きの行を上から積み、名前→セル番地を覚えておく。数式は {名前} で他のセルを参照する。"""

    def __init__(self, wb, title, names, widths):
        self.ws = wb.create_sheet(title)
        self.title = title
        self.names = names
        self.r = 1
        for i, w in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = w

    def f(self, expr):
        # {名前} は保存の直前にまとめて番地へ置き換える（後のシートで決まる名前も参照できるように）
        return "=" + expr

    def name(self, key, col, row=None):
        assert key not in self.names, f"名前の重複: {key}"
        self.names[key] = (self.title, f"{get_column_letter(col)}{row or self.r}")

    def cell(self, col, value, font=F_CALC, fmt=None, fill=None, border=True, align=None):
        c = self.ws.cell(row=self.r, column=col, value=value)
        c.font = font
        if fmt:
            c.number_format = fmt
        if fill:
            c.fill = fill
        if border:
            c.border = BOX
        c.alignment = align or Alignment(vertical="center", wrap_text=isinstance(value, str) and not str(value).startswith("="))
        return c

    def title_row(self, text, note=None):
        self.ws.cell(row=self.r, column=1, value=text).font = F_TITLE
        self.r += 1
        if note:
            self.ws.cell(row=self.r, column=1, value=note).font = F_NOTE
            self.r += 1
        self.r += 1

    def section(self, text, heads=None):
        self.r += 1 if self.r > 3 else 0
        self.ws.cell(row=self.r, column=1, value=text).font = F_SEC
        self.r += 1
        if heads:
            for i, h in enumerate(heads, 1):
                self.cell(i, h, F_HEAD, fill=BG_HEAD, align=Alignment(horizontal="center", vertical="center", wrap_text=True))
            self.r += 1

    def note(self, text, col=1):
        self.ws.cell(row=self.r, column=col, value=text).font = F_NOTE
        self.r += 1


def _ref(names, name, here):
    sheet, addr = names[name]
    return addr if sheet == here else f"'{sheet}'!{addr}"


def build(ed, R, out_dir):
    s, b = ed["sample"], ed["brand"]
    c, props = s["customer"], s["props"]
    names = {}
    wb = Workbook()
    wb.remove(wb.active)

    # ───────── ① かんたん試算（1枚で完結：入力 → 結果 → チェック → まとめ → グラフ） ─────────
    e = Sheet(wb, S_EASY, names, [30, 22, 24, 24, 56])
    WRAP = Alignment(wrap_text=True, vertical="center")
    e.ws.cell(row=1, column=1, value=f"資金計画シミュレーター　{b['name']}").font = Font(name="Arial", size=16, bold=True, color=NAVY)
    e.ws.cell(row=2, column=1, value="人生を、自分の土地から。　単身者向け 戸建ての資金計画（概算）").font = F_NOTE
    e.ws.cell(row=3, column=1, value="使い方：① 黄色のマスに入力する　→　② 下の「結果」と「チェック」を見る　→　③ もっと詳しく見るときは ②〜⑥ のタブ").font = Font(name="Arial", size=11, bold=True, color=GOLD)
    e.r = 4
    dv_kind = DataValidation(type="list", formula1='"新築,買取再販,中古"', allow_blank=False)
    e.ws.add_data_validation(dv_kind)

    def ein(key, label, value, fmt, hint, cols=(2,), dv=None):
        e.cell(1, label)
        for k, col in enumerate(cols):
            cl = e.cell(col, value[k] if isinstance(value, list) else value, F_IN, fmt, BG_IN)
            if dv:
                dv.add(cl)
            e.name(key.replace("#", "AB"[k]) if "#" in key else key, col)
        e.cell(5, hint, F_NOTE)
        e.r += 1

    e.section("STEP 1　あなたのこと", ["項目", "入力", "", "", "入力のヒント"])
    ein("e_gross", "年収（税込）", round(c["gross"] / 10_000), MAN, "源泉徴収票の「支払金額」。万円で入力（450万円なら 450）")
    ein("e_rent", "今の家賃（月）", c["rent"], YEN, "管理費・共益費を含む")
    ein("e_park", "駐車場代（月）", c.get("rentParking", 0), YEN, "車を置いていなければ 0")
    ein("e_sav", "預貯金", round(c["savings"] / 10_000), MAN, "すぐに使える預金の合計。万円で入力")
    ein("e_exist", "車などのローンの返済（月）", c.get("existingMonthly", 0), YEN, "なければ 0。ローン審査では住宅ローンと合わせて見られます")
    e.section("STEP 2　住宅ローン", ["項目", "入力", "", "", "入力のヒント"])
    ein("e_rate", "金利（年）", s["loan"]["rate"], PCT2, "全期間この金利と仮定。変動金利なら今の金利")
    ein("e_years", "返済年数", s["loan"]["years"], YEARS, "35年が多い")
    e.section("STEP 3　くらべる物件", ["項目", "", "物件A", "物件B", "入力のヒント"])
    ein("e_label#", "物件の名前", [p["label"] for p in props], None, "わかりやすい呼び名", (3, 4))
    ein("e_price#", "物件価格", [round(p["price"] / 10_000) for p in props], MAN, "販売価格。万円で入力", (3, 4))
    ein("e_down#", "頭金", [round(p["down"] / 10_000) for p in props], MAN, "0でも試算できます。万円で入力", (3, 4))
    ein("e_kind#", "区分", [KIND[p["kind"]] for p in props], None, "新築／買取再販（リフォーム済みの中古）／中古。住宅ローン控除の判定に使う", (3, 4), dv_kind)

    e.section("結果（自動で計算・概算）", ["項目", "今の賃貸", "物件A", "物件B", "この数字の意味"])
    head_row = e.r - 1
    DASH = Alignment(horizontal="center", vertical="center")

    def eout(key, label, rent_expr, expr, fmt, meaning, key_row=False):
        font, fill = (F_KEY, BG_KEY) if key_row else (F_CALC, None)
        e.cell(1, label, font, fill=fill)
        if rent_expr is None:
            e.cell(2, "—", F_NOTE, align=DASH)
        else:
            e.cell(2, e.f(rent_expr), font, fmt, fill)
        for k, L in enumerate("AB"):
            e.name(f"{key}{L}", 3 + k)
            e.cell(3 + k, e.f(expr.replace("#", L)), font, fmt, fill, align=WRAP if fmt is None else None)
        e.cell(5, meaning, F_NOTE)
        e.r += 1

    eout("eNm", "物件の名前", '"今の住まい"', "{label#}", None, "")
    mo_row = e.r
    eout("eMo", "毎月の住まいの支払い", "{rmA}", "{om#}", YEN, "ローン返済＋固定資産税・保険・修繕の積立。賃貸は家賃＋更新料・保険・駐車場の月割り", True)
    eout("ePm", "うちローンの返済", None, "{pm#}", YEN, "毎月同じ額を返す方式（元利均等返済）")
    eout("eDf", "今の家賃との差", None, "{dm#}", SIGNED, "＋は今より増える、−は今より減る")
    eout("eCo", "はじめに必要な現金", None, "{co#}/10000", MAN1, "頭金＋現金で払う諸費用（諸費用もローンにするときは頭金だけ）")
    eout("eCl", "手元に残る預貯金", "{savings}/10000", "{cl#}/10000", MAN1, "買ったあとに残る預金")
    eout("eCr", "住宅ローン控除（最初の10年）", None, "{c10#}/10000", MAN1, "税金から戻る見込み額。税額が少ないと全額は戻りません")
    eout("eTn", "10年間の住まいの費用", "{t7A}/10000", "{t6#}/10000", MAN1, "10年分の支払い＋はじめの現金−控除。賃貸は今の家賃が続く場合", True)

    e.section("チェック（◎ 安心　△ 注意　× 要相談）", ["チェック", "基準", "物件A", "物件B", "見かた"])
    chk_top = e.r

    def echk(label, crit, expr, meaning):
        e.cell(1, label, F_KEY)
        e.cell(2, e.f(crit), F_NOTE, align=WRAP)
        for k, L in enumerate("AB"):
            e.cell(3 + k, e.f(expr.replace("#", L)), F_KEY, align=WRAP)
        e.cell(5, meaning, F_NOTE)
        e.ws.row_dimensions[e.r].height = 36
        e.r += 1

    echk("返済の負担", '"◎ 年収の"&TEXT({comfort},"0%")&"以内"',
         'IF({ratio#}<={comfort},"◎ 目安内",IF({ratio#}<={sl#},"△ やや重い","× 重い"))&"（"&TEXT({ratio#},"0.0%")&"）"',
         "ローン返済＋車などの返済が、年収の何%か。△は審査の基準以内だが余裕は少ない")
    echk("ローン審査の目安", '"審査金利"&TEXT({screenRate},"0.0%")&"で"&TEXT({slA},"0%")&"以内"',
         'IF({sr#}<={sl#},"◎ 基準内","× 基準を超える")&"（"&TEXT({sr#},"0.0%")&"）"',
         "金融機関は高めの金利で返済の割合を見ます。実際の審査は金融機関ごとに違います")
    echk("もしもの貯金", '"住まい＋生活費の6か月分"',
         'IF({cl#}>={rs#},"◎ 6か月分以上",IF({cl#}>={rs#}/2,"△ 3か月分以上","× 3か月分未満"))',
         "病気や転職のときに、手元の預金で何か月暮らせるか")
    echk("住宅ローン控除", '"単身者の一般の限度額で判定"',
         'IF({ok#}="はい","◎ 対象（"&{yrs#}&"年間）","× "&{why#})',
         "×のときは、最初に満たさない要件を表示。詳しくは ③ 詳しい試算")
    rng = f"C{chk_top}:D{e.r - 1}"
    for mark, fill in (("◎", BG_OK), ("△", BG_WARN), ("×", BG_NG)):
        e.ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(C{chk_top},1)="{mark}"'], fill=fill))

    e.section("ひとことまとめ")
    for L in "AB":
        e.cell(1, f"物件{L}", F_KEY)
        e.ws.merge_cells(start_row=e.r, start_column=2, end_row=e.r, end_column=5)
        e.cell(2, e.f('"「"&{label' + L + '}&"」は、毎月 約"&TEXT({om' + L + '}/10000,"0.0")&"万円（今より"&IF({dm' + L + '}>=0,"＋","−")&TEXT(ABS({dm' + L + '})/10000,"0.0")&"万円）。'
                      'はじめに現金 "&TEXT({co' + L + '}/10000,"0")&"万円、住宅ローン控除は最初の10年で約"&TEXT({c10' + L + '}/10000,"0")&"万円の見込みです。"'),
               Font(name="Arial", size=11, color="000000"), align=WRAP)
        e.ws.row_dimensions[e.r].height = 36
        e.r += 1

    e.section("ご注意")
    for t in [
        "・概算です。実際の返済額・審査結果・税額・控除額は、金融機関や税務署の判断、ご本人の所得で変わります。",
        "・住宅ローン控除は単身者向けの一般の限度額で計算し、控除しきれない額は含めていません。〔要確認〕の項目は「参考：控除の制度」に記載。",
        f"・制度は{R['BASIS_DATE']}時点（2026〜2030年入居）。出典：" + "、".join(R["SOURCES"]) + "。",
        f"{b['name']}　TEL {b['tel']}　{b['site']}" + (f"　{b['address']}" if b.get("address") else "") + f"　宅地建物取引業 免許番号：{b['license']}",
    ]:
        e.ws.merge_cells(start_row=e.r, start_column=1, end_row=e.r, end_column=5)
        e.ws.cell(row=e.r, column=1, value=t).font = F_NOTE
        e.ws.cell(row=e.r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        e.r += 1

    ch = BarChart()
    ch.type = "col"
    ch.title = "毎月の住まいの支払い"
    ch.legend = None
    ch.y_axis.numFmt = '#,##0'
    ch.add_data(Reference(e.ws, min_col=2, max_col=4, min_row=mo_row, max_row=mo_row), from_rows=True, titles_from_data=False)
    ch.set_categories(Reference(e.ws, min_col=2, max_col=4, min_row=head_row, max_row=head_row))
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    for k_ in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
        setattr(ch.dataLabels, k_, False)
    ch.height, ch.width = 7.5, 12
    e.ws.add_chart(ch, "G4")

    # ───────── 入力 ─────────
    i = Sheet(wb, S_IN, names, [34, 22, 22, 46])
    i.title_row("② 詳しい条件", "ふだんは変えなくてOKな既定値です。緑の文字は ① かんたん試算 で入力した値（ここでは変えられません）。")
    LINKED = "① かんたん試算で入力"

    def inp(key, label, value, fmt=YEN, note="", col=2, link=None):
        i.cell(1, label)
        if link:
            i.cell(col, i.f(link), F_LINK, fmt)
            note = LINKED
        else:
            i.cell(col, value, F_IN, fmt, BG_IN)
        if note:
            i.cell(4, note, F_NOTE)
        i.name(key, col)
        i.r += 1

    i.section("お客様", ["項目", "値", "", "メモ"])
    inp("gross", "年収（税込）", None, link="{e_gross}*10000")
    inp("social", "社会保険料（年収に対する割合）", c.get("socialRate", 0.15), PCT, "概算。源泉徴収票の社会保険料等の金額÷年収")
    inp("rent", "家賃（月額）", None, link="{e_rent}")
    inp("rentPark", "駐車場代（月額）", None, link="{e_park}")
    inp("rentIns", "賃貸の火災保険（年額）", c.get("rentInsYear", 20000))
    inp("renewM", "更新料（家賃の何か月分）", 1, NUM)
    inp("renewEvery", "更新の間隔（か月）", 24, '0')
    inp("savings", "預貯金", None, link="{e_sav}*10000")
    inp("living", "生活費（月額・住居費を除く）", c["livingMonth"])
    inp("existing", "他の借入の返済（月額）", None, link="{e_exist}")
    inp("wish", "毎月の支払いの希望額", c.get("monthlyWish", 0), note="0のときは「家賃並みで買える価格」を出さない")
    inp("moveYear", "入居年", c["moveInYear"], '0')
    inp("moveMonth", "入居月", c["moveInMonth"], '0', "この月に借入を実行し、翌月から返済すると仮定")

    i.section("住宅ローン", ["項目", "値", "", "メモ"])
    inp("rate", "金利（年）", None, PCT2, link="{e_rate}")
    inp("years", "返済年数", None, YEARS, link="{e_years}")
    inp("screenRate", "審査金利（年）", (s.get("screen") or {}).get("rate", 0.03), PCT2, "金融機関が審査で使う金利の目安（仮定）")
    inp("limLow", "審査の返済比率の基準（年収が境目未満）", 0.30, PCT, "フラット35の基準。金融機関で異なる")
    inp("limHigh", "審査の返済比率の基準（年収が境目以上）", 0.35, PCT)
    inp("limLine", "審査基準の年収の境目", 4_000_000)
    inp("comfort", "無理のない返済の目安（年収比）", ed.get("comfortRatio", 0.25), PCT, "既存の借入の返済を含む")

    i.section("物件", ["項目", "物件A", "物件B", "メモ"])

    def pin(key, label, vals, fmt=YEN, note="", dv=None, link=None, calc=False):
        i.cell(1, label)
        for k in range(2):
            L = "AB"[k]
            if link:
                cl = i.cell(2 + k, i.f(link.replace("#", L)), F_CALC if calc else F_LINK, fmt)
            else:
                cl = i.cell(2 + k, vals[k], F_IN, fmt, BG_IN)
            if dv:
                dv.add(cl)
            i.name(f"{key}{L}", 2 + k)
        if link and not calc:
            note = LINKED
        if note:
            i.cell(4, note, F_NOTE)
        i.r += 1

    dv_perf = DataValidation(type="list", formula1='"' + ",".join(PERF.values()) + '"', allow_blank=False)
    dv_yn = DataValidation(type="list", formula1='"はい,いいえ"', allow_blank=False)
    for v in (dv_perf, dv_yn):
        i.ws.add_data_validation(v)
    P = props
    pin("label", "名称", None, None, link="{e_label#}")
    pin("kind", "区分", None, None, link="{e_kind#}")
    pin("perf", "省エネ性能", [PERF[p["perf"]] for p in P], None, "証明書で確認できる性能", dv_perf)
    pin("cert", "性能の証明書がある", ["はい" if p["certs"] else "いいえ" for p in P], None, "ないときは「その他」として判定", dv_yn)
    pin("price", "物件価格", None, link="{e_price#}*10000")
    pin("landShare", "物件価格のうち土地の割合", [p["landPrice"] / p["price"] for p in P], PCT, "売却の目安に使う（残りが建物）")
    pin("land", "うち土地", None, link="{price#}*{landShare#}", calc=True)
    pin("bldg", "うち建物", None, link="{price#}-{land#}", calc=True)
    pin("floor", "床面積（㎡）", [p["floor"] for p in P], NUM)
    pin("built", "建築日（新築は完成予定日）", [d(p["builtDate"]) for p in P], DATE)
    pin("confirm", "建築確認日", [d(p["confirmDate"]) for p in P], DATE, "わからなければ空欄")
    pin("down", "頭金", None, link="{e_down#}*10000")
    pin("costRate", "諸費用（物件価格に対する割合）", [p["buyCostRate"] for p in P], PCT)
    pin("fin", "諸費用もローンで借りる", ["はい" if p.get("financeCosts") else "いいえ" for p in P], None, "", dv_yn)
    pin("propTax", "固定資産税・都市計画税（年額）", [p["propTaxYear"] for p in P])
    pin("ins", "火災・地震保険（年額）", [p["insYear"] for p in P])
    pin("repair", "修繕の積立（月額）", [p["repairMonth"] for p in P], note="自分で積み立てる額の目安")
    pin("b10", "10年後の建物価値（今の価値に対する割合）", [p["b10"] for p in P], PCT, "売却の目安に使う仮定")
    pin("relative", "親族などからの取得", ["はい" if p.get("fromRelative") else "いいえ" for p in P], None, "「はい」は控除の対象外", dv_yn)
    pin("within6", "取得から6か月以内に入居する", ["いいえ" if p.get("moveInOK") is False else "はい" for p in P], None, "", dv_yn)
    pin("seismic", "耐震基準の証明がある（1981年以前の中古）", ["はい" if p.get("seismicCert") else "いいえ" for p in P], None, "", dv_yn)
    pin("red", "災害レッドゾーン内", ["はい" if p.get("redZone") else "いいえ" for p in P], None, "2028年以降入居の新築は対象外", dv_yn)

    # ───────── 税の前提（表） ─────────
    t = Sheet(wb, S_TAXT, names, [26, 16, 16, 40])
    t.title_row("参考：税の前提", "令和8・9年分の所得税の特例を2028年以降も続くと仮定（税額＝控除できる余地を低めに見積もる）。")

    def table(key, title, heads, rows, fmts, note=None):
        t.section(title, heads)
        top = t.r
        for row in rows:
            for k, v in enumerate(row):
                t.cell(1 + k, v, F_IN if not isinstance(v, str) else F_CALC, fmts[k])
            t.r += 1
        names[key] = (t.title, f"$A${top}:${get_column_letter(len(heads))}${t.r - 1}")
        if note:
            t.note(note)

    table("tblSalary", "給与所得控除（給与収入がこの額以上のとき：収入×率＋加算額）", ["給与収入（以上）", "率", "加算額"],
          [[0, 0, 550_000], [1_625_001, 0.4, -100_000], [1_800_001, 0.3, 80_000], [3_600_001, 0.2, 440_000],
           [6_600_001, 0.1, 1_100_000], [8_500_001, 0, 1_950_000]], [YEN, PCT, YEN])
    t.cell(1, "給与所得控除の最低額")
    t.cell(2, 740_000, F_IN, YEN, BG_IN)
    t.cell(4, "令和8・9年分の特例（74万円）", F_NOTE)
    t.name("salMin", 2)
    t.r += 1
    table("tblBasic", "所得税の基礎控除（合計所得金額がこの額を超えるとき）", ["合計所得（超）", "基礎控除"],
          [[0, 1_040_000]] + [[v + 0.001, d_] for v, d_ in [(4_890_000, 670_000), (6_550_000, 620_000), (23_500_000, 480_000),
                                                       (24_000_000, 320_000), (24_500_000, 160_000), (25_000_000, 0)]],
          [YEN, YEN], "「超える」を表すため、境目には0.001円を足しています（表示は円単位）。")
    table("tblIT", "所得税の税率（課税所得がこの額を超えるとき：課税所得×税率−控除額）", ["課税所得（以上）", "税率", "控除額"],
          [[0, 0.05, 0], [1_951_000, 0.1, 97_500], [3_301_000, 0.2, 427_500], [6_951_000, 0.23, 636_000],
           [9_001_000, 0.33, 1_536_000], [18_001_000, 0.4, 2_796_000], [40_001_000, 0.45, 4_796_000]], [YEN, PCT, YEN])
    t.section("その他の定数", ["項目", "値", "", "メモ"])
    for key, label, v, fmt, note in [
        ("rtBasic", "住民税の基礎控除", 430_000, YEN, ""),
        ("rtRate", "住民税の所得割の税率", 0.10, PCT, ""),
        ("rtAdj", "住民税の調整控除（概算）", 2_500, YEN, "単身の概算"),
        ("rtFlat", "住民税の均等割など（年額）", 5_000, YEN, "手取りの概算に使う"),
        ("recon", "復興特別所得税", 0.021, PCT, "所得税×2.1%"),
    ]:
        t.cell(1, label)
        t.cell(2, v, F_IN, fmt, BG_IN)
        if note:
            t.cell(4, note, F_NOTE)
        t.name(key, 2)
        t.r += 1

    # ───────── 控除の制度 ─────────
    rr = Sheet(wb, S_RULE, names, [12, 22, 10, 10, 16, 10, 14, 14, 60])
    rr.title_row("住宅ローン控除の制度（単身者・一般の借入限度額）",
                 f"{R['BASIS_DATE']}時点。税制改正があったら、この表を書き換えます。")
    rr.section("区分ごとの借入限度額と控除期間", ["区分", "省エネ性能", "入居年から", "入居年まで", "借入限度額", "控除期間（年）",
                                         "建築確認の期限", "建築の期限", "注記"])
    top = rr.r
    for rule in R["RULES"]:
        for k in rule["kinds"]:
            cond = rule.get("cond") or {}
            note = rule.get("note", "")
            if rule.get("check") or k == "resale":
                note = (note + "　" if note else "") + "買取再販の限度額は新築に準じる〔要確認〕"
            row = [KIND[k], PERF[rule["perf"]], rule["from"], rule["to"], rule["limit"][0], rule["years"],
                   d(cond.get("confirmBy")), d(cond.get("builtBy")), note]
            for n, v in enumerate(row):
                rr.cell(1 + n, v, F_IN if n in (2, 3, 4, 5, 6, 7) else F_CALC, [None, None, '0', '0', YEN, '0', DATE, DATE, None][n])
            rr.r += 1
    bot = rr.r - 1
    for key, col in [("rK", "A"), ("rP", "B"), ("rFrom", "C"), ("rTo", "D"), ("rLim", "E"), ("rYears", "F"), ("rConf", "G"), ("rBuilt", "H")]:
        names[key] = (rr.title, f"${col}${top}:${col}${bot}")
    rr.note("建築確認の期限・建築の期限がある区分は、どちらかの期限までに満たしたものに限り対象。")
    rr.section("共通の要件", ["項目", "値"])
    for key, label, v, fmt in [
        ("cRate", "控除率", 0.007, PCT),
        ("cFloor", "床面積の下限（㎡）", 40, NUM),
        ("cFloorFull", "この床面積未満は合計所得1,000万円以下の年だけ", 50, NUM),
        ("cIncSmall", "床面積が小さい場合の合計所得の上限", 10_000_000, YEN),
        ("cInc", "合計所得金額の上限", 20_000_000, YEN),
        ("cYears", "返済期間の下限（年）", 10, '0'),
        ("cSeismic", "中古：この日以降の建築なら耐震の証明不要", date(1982, 1, 1), DATE),
        ("cRed", "この入居年以降、レッドゾーンの新築は対象外", 2028, '0'),
        ("cResRate", "住民税から控除できる上限（課税総所得金額等に対する割合）", 0.05, PCT),
        ("cResMax", "住民税から控除できる上限額〔要確認〕", 97_500, YEN),
    ]:
        rr.ws.merge_cells(start_row=rr.r, start_column=1, end_row=rr.r, end_column=4)
        rr.cell(1, label)
        rr.cell(5, v, F_IN, fmt, BG_IN)
        rr.name(key, 5)
        rr.r += 1

    # ───────── 税額の概算 ─────────
    x = Sheet(wb, S_TAX, names, [34, 16, 16, 16, 40])
    x.title_row("税額の概算（給与収入だけの単身者）", "住宅ローン控除と「収入が減ったら」に使います。源泉徴収票がある場合は数字を見比べてください。")
    x.section("年収の変化", ["項目", "今の年収", "2割減", "4割減", "メモ"])
    x.cell(1, "年収の減り方")
    for k, v in enumerate([0, 0.2, 0.4]):
        x.cell(2 + k, v, F_IN, PCT, BG_IN)
        x.name(f"drop{k}", 2 + k)
    x.r += 1
    rows = [
        ("g", "年収", "{gross}*(1-{drop@})", YEN),
        ("sd", "給与所得控除", "MIN({g@},MAX({salMin},{g@}*VLOOKUP({g@},{tblSalary},2,TRUE)+VLOOKUP({g@},{tblSalary},3,TRUE)))", YEN),
        ("ti", "合計所得金額", "MAX(0,{g@}-{sd@})", YEN),
        ("so", "社会保険料", "{g@}*{social}", YEN),
        ("bd", "所得税の基礎控除", "VLOOKUP({ti@},{tblBasic},2,TRUE)", YEN),
        ("tx", "課税所得", "INT(MAX(0,{ti@}-{so@}-{bd@})/1000)*1000", YEN),
        ("it", "所得税（控除前）", "MAX(0,INT({tx@}*VLOOKUP({tx@},{tblIT},2,TRUE)-VLOOKUP({tx@},{tblIT},3,TRUE)))", YEN),
        ("rtx", "住民税の課税所得", "INT(MAX(0,{ti@}-{so@}-{rtBasic})/1000)*1000", YEN),
        ("rt", "住民税（所得割）", "MAX(0,INT({rtx@}*{rtRate}-{rtAdj}))", YEN),
        ("th", "手取り（年額・概算）", "{g@}-{so@}-{it@}*(1+{recon})-{rt@}-{rtFlat}", YEN),
        ("thm", "手取り（月額・概算）", "{th@}/12", YEN),
    ]
    for key, label, expr, fmt in rows:
        x.cell(1, label)
        for k in range(3):
            x.name(f"{key}{k}", 2 + k)
            x.cell(2 + k, x.f(expr.replace("@", str(k))), F_LINK if key == "g" else F_CALC, fmt)
        x.r += 1

    # ───────── 試算 ─────────
    m = Sheet(wb, S_CALC, names, [40, 20, 20, 52])
    m.title_row("③ 詳しい試算", "① かんたん試算・② 詳しい条件 から自動で計算します（概算）。グレーの「計算用」の行は ＋ で開けます。")

    def row(key, label, expr, fmt=YEN, note="", key_row=False, font=None):
        m.cell(1, label, F_KEY if key_row else F_CALC, fill=BG_KEY if key_row else None)
        for k, L in enumerate("AB"):
            m.name(f"{key}{L}", 2 + k)
            m.cell(2 + k, m.f(expr.replace("#", L)), font or (F_KEY if key_row else F_CALC), fmt, BG_KEY if key_row else None)
        if note:
            m.cell(4, note, F_NOTE)
        m.r += 1

    m.section("物件", ["項目", "物件A", "物件B", "説明"])
    row("nm", "名称", "{label#}", None, font=F_LINK)
    row("pr", "物件価格", "{price#}", font=F_LINK)
    m.section("借入と初期費用", ["項目", "物件A", "物件B", "説明"])
    row("bc", "諸費用", "{price#}*{costRate#}")
    row("el", "住宅の取得に充てる借入", "MAX(0,{price#}-{down#})", note="住宅ローン控除の対象になる部分")
    row("ln", "借入額", '{el#}+IF({fin#}="はい",{bc#},0)', key_row=True, note="諸費用もローンで借りる場合は諸費用を含む")
    row("pm", "毎月の返済額", "IF({ln#}<=0,0,IF({rate}=0,{ln#}/({years}*12),PMT({rate}/12,{years}*12,-{ln#})))", key_row=True, note="元利均等返済")
    row("co", "初期費用（現金）", '{down#}+IF({fin#}="はい",0,{bc#})', note="頭金＋現金で払う諸費用")
    row("cl", "手元に残る預貯金", "{savings}-{co#}")
    m.section("毎月の住居費", ["項目", "物件A", "物件B", "説明"])
    row("om", "持ち家の住居費（月額）", "{pm#}+{propTax#}/12+{ins#}/12+{repair#}", key_row=True, note="返済＋税金＋保険＋修繕の積立")
    row("upk", "うち返済以外", "{om#}-{pm#}")
    row("rm", "今の賃貸の住居費（月額）", "{rent}+{rent}*{renewM}/{renewEvery}+{rentIns}/12+{rentPark}", note="家賃＋更新料の月割り＋保険＋駐車場")
    row("dm", "賃貸との差（持ち家−賃貸）", "{om#}-{rm#}")
    row("rs", "6か月分の予備資金の目安", "({om#}+{living})*6", note="住居費＋生活費の6か月分")
    row("rsok", "手元資金は予備資金の目安以上か", 'IF({cl#}>={rs#},"はい","不足")', None)
    m.section("返済の負担と審査の目安", ["項目", "物件A", "物件B", "説明"])
    row("ratio", "返済比率（年収比・既存の借入を含む）", "IF({gross}=0,0,({pm#}*12+{existing}*12)/{gross})", PCT, key_row=True)
    row("rok", "無理のない返済の目安以内か", 'IF({ratio#}<={comfort},"はい","超える")', None)
    row("thr", "参考：手取りに対する返済の割合", "IF({thm0}=0,0,{pm#}/{thm0})", PCT, "手取りは税額の概算シートの今の年収")
    row("sp", "審査金利での毎月の返済額", "IF({ln#}<=0,0,IF({screenRate}=0,{ln#}/({years}*12),PMT({screenRate}/12,{years}*12,-{ln#})))")
    row("sr", "審査の返済比率", "IF({gross}=0,0,({sp#}*12+{existing}*12)/{gross})", PCT, "審査金利で計算した返済＋既存の借入")
    row("sl", "審査の基準", "IF({gross}<{limLine},{limLow},{limHigh})", PCT)
    row("sok", "審査の基準以内か（目安）", 'IF({sr#}<={sl#},"はい","超える")', None, "実際の審査は金融機関ごとに異なる")
    row("room", "住宅ローンに使える毎月の上限（審査）", "MAX(0,{gross}*{sl#}/12-{existing})")
    row("mx", "審査上の借入の上限（目安）", "IF({screenRate}=0,{room#}*{years}*12,PV({screenRate}/12,{years}*12,-{room#}))")
    m.section("住宅ローン控除の判定", ["項目", "物件A", "物件B", "説明"])
    helper_top = m.r
    row("jp", "（計算用）判定に使う省エネ性能", 'IF(AND({perf#}<>"その他",{cert#}<>"はい"),"その他",{perf#})', None, "証明書がない場合は「その他の住宅」")
    row("cnt", "（計算用）制度の該当行の数", "COUNTIFS({rK},{kind#},{rP},{jp#},{rFrom},\"<=\"&{moveYear},{rTo},\">=\"&{moveYear})", '0')
    row("cfb", "（計算用）建築確認の期限（経過措置）", "SUMIFS({rConf},{rK},{kind#},{rP},{jp#},{rFrom},\"<=\"&{moveYear},{rTo},\">=\"&{moveYear})", DATE0)
    row("bfb", "（計算用）建築の期限（経過措置）", "SUMIFS({rBuilt},{rK},{kind#},{rP},{jp#},{rFrom},\"<=\"&{moveYear},{rTo},\">=\"&{moveYear})", DATE0)
    m.ws.row_dimensions.group(helper_top, m.r - 1, outline_level=1, hidden=True)
    checks = [
        ("k1", "入居年が制度の対象（2026〜2030年）", "AND({moveYear}>=MIN({rFrom}),{moveYear}<=MAX({rTo}))"),
        ("k2", "親族などからの取得ではない", '{relative#}<>"はい"'),
        ("k3", "返済期間が10年以上", "{years}>={cYears}"),
        ("k4", "床面積が40㎡以上", "{floor#}>={cFloor}"),
        ("k5", "取得から6か月以内に入居", '{within6#}="はい"'),
        ("k6", "耐震（1981年以前の中古は証明が必要）", 'NOT(AND({kind#}="中古",ISNUMBER({built#}),{built#}<{cSeismic},{seismic#}<>"はい"))'),
        ("k7", "災害レッドゾーンの新築ではない（2028年以降）", 'NOT(AND({kind#}="新築",{red#}="はい",{moveYear}>={cRed}))'),
        ("k8", "区分・省エネ性能が制度の対象", "{cnt#}>0"),
        ("k9", "経過措置の期限を満たす", "IF({cfb#}=0,TRUE,OR(AND(ISNUMBER({confirm#}),{confirm#}<={cfb#}),AND(ISNUMBER({built#}),{built#}<={bfb#})))"),
        ("k10", "所得の要件（合計所得2,000万円以下など）", "AND({ti0}<={cInc},OR({floor#}>={cFloorFull},{ti0}<={cIncSmall}))"),
    ]
    first_check = m.r
    for key, label, expr in checks:
        row(key, label, f'IF({expr},"○","×")', None)
    last_check = m.r - 1
    for L, col in (("A", "B"), ("B", "C")):
        names[f"chk{L}"] = (S_CALC, f"{col}{first_check}:{col}{last_check}")
    names["chkLabels"] = (S_CALC, f"$A${first_check}:$A${last_check}")
    row("ok", "控除の対象か", 'IF(COUNTIF({chk#},"×")=0,"はい","いいえ")', None, key_row=True, note="上の要件がすべて ○ なら対象")
    row("why", "判定", 'IF({ok#}="はい","対象："&{kind#}&"・"&{jp#},"対象外："&INDEX({chkLabels},MATCH("×",{chk#},0)))', None,
        note="対象外のときは、最初に満たさない要件")
    row("lim", "借入限度額", "IF({ok#}=\"はい\",SUMIFS({rLim},{rK},{kind#},{rP},{jp#},{rFrom},\"<=\"&{moveYear},{rTo},\">=\"&{moveYear}),0)")
    row("yrs", "控除期間（年）", "IF({ok#}=\"はい\",SUMIFS({rYears},{rK},{kind#},{rP},{jp#},{rFrom},\"<=\"&{moveYear},{rTo},\">=\"&{moveYear}),0)", '0')
    row("cmax", "制度上の控除の上限（期間合計）", "SUM({cMax#})", note="年末残高×0.7%の合計。税額が足りないと全額は控除できない")
    row("ctot", "控除見込額（期間合計・概算）", "SUM({cAct#})", key_row=True, note="所得税→住民税の順に控除。控除しきれない額は含めない")
    row("c1", "初年の控除見込額", "{cAct1#}")
    row("c10", "最初の10年の控除見込額", "SUM({cAct10#})")
    m.section("10年間の住まいの費用", ["項目", "物件A", "物件B", "説明"])
    row("t1", "初期費用（現金）", "{co#}")
    row("t2", "返済（10年分）", "{pm#}*120")
    row("t3", "税金・保険（10年分）", "({propTax#}+{ins#})*10")
    row("t4", "修繕の積立（10年分）", "{repair#}*120")
    row("t5", "控除見込額（10年分）", "-{c10#}")
    row("t6", "10年間の合計（控除あり）", "SUM({t1#},{t2#},{t3#},{t4#},{t5#})", key_row=True)
    row("t7", "参考：同じ10年の賃貸の住居費", "{rm#}*120", note="今の家賃が続く場合")
    m.section("家賃並みで買える価格（物件Aの頭金・諸費用の条件）", ["項目", "値", "", "説明"])
    m.cell(1, "借入できる額")
    m.cell(2, m.f("IF({wish}<=0,0,IF({rate}=0,{wish}*{years}*12,PV({rate}/12,{years}*12,-{wish})))"))
    m.ws.cell(row=m.r, column=2).number_format = YEN
    m.name("afLoan", 2)
    m.cell(4, "毎月の支払いの希望額を返済に充てる場合", F_NOTE)
    m.r += 1
    m.cell(1, "買える物件価格の目安", F_KEY, fill=BG_KEY)
    m.cell(2, m.f('IF({wish}<=0,0,IF({finA}="はい",({afLoan}+{downA})/(1+{costRateA}),{afLoan}+{downA}))'), F_KEY, YEN, BG_KEY)
    m.name("afPrice", 2)
    m.cell(4, "希望額が0のときは0", F_NOTE)
    m.r += 1
    m.cell(1, "必要な現金")
    m.cell(2, m.f('IF({wish}<=0,0,{downA}+IF({finA}="はい",0,{afPrice}*{costRateA}))'), fmt=YEN)
    m.r += 1

    # ───────── 住宅ローン控除（年ごと） ─────────
    y = Sheet(wb, S_CR, names, [8, 8, 15, 15, 15, 15, 15, 15, 15])
    y.title_row("④ 住宅ローン控除の見込み（年ごと）",
                "上限＝min（年末残高×住宅取得分の割合, 借入限度額, 物件価格）×0.7%（100円未満切捨て）。所得税→住民税の順に控除し、控除しきれない額は加算しない。")
    heads = ["年目", "年", "年末残高", "控除対象の残高", "控除の上限", "所得税から", "住民税から", "控除しきれない額", "控除見込額"]
    for L in "AB":
        y.section(f"物件{L}", heads)
        y.ws.cell(row=y.r - 2, column=3).value = y.f(f"{{label{L}}}")
        y.ws.cell(row=y.r - 2, column=3).font = F_LINK
        top = y.r
        for n in range(1, 14):
            r_ = y.r
            ln, cell = f"{{ln{L}}}", lambda col: f"{col}{r_}"
            k = f"MIN(12-{{moveMonth}}+12*(A{r_}-1),{{years}}*12)"
            bal = (f"IF({{rate}}=0,MAX(0,{ln}-{{pm{L}}}*{k}),MAX(0,{ln}*(1+{{rate}}/12)^{k}-{{pm{L}}}*((1+{{rate}}/12)^{k}-1)/({{rate}}/12)))")
            vals = [
                (n, '0', F_CALC),
                (y.f(f"{{moveYear}}+A{r_}-1"), '0', F_CALC),
                (y.f(bal), YEN, F_CALC),
                (y.f(f"IF({ln}<=0,0,C{r_}*MIN(1,{{el{L}}}/{ln}))"), YEN, F_CALC),
                (y.f(f"IF(AND({{ok{L}}}=\"はい\",A{r_}<={{yrs{L}}}),INT(MIN(D{r_},{{lim{L}}},{{price{L}}})*{{cRate}}/100)*100,0)"), YEN, F_CALC),
                (y.f(f"MIN(E{r_},{{it0}})"), YEN, F_CALC),
                (y.f(f"MIN(E{r_}-F{r_},MIN({{cResMax}},INT({{tx0}}*{{cResRate}}),{{rt0}}))"), YEN, F_CALC),
                (y.f(f"E{r_}-F{r_}-G{r_}"), YEN, F_CALC),
                (y.f(f"F{r_}+G{r_}"), YEN, F_KEY),
            ]
            for col, (v, fmt, font) in enumerate(vals, 1):
                y.cell(col, v, font, fmt)
            y.r += 1
        bot = y.r - 1
        names[f"cMax{L}"] = (S_CR, f"E{top}:E{bot}")
        names[f"cAct{L}"] = (S_CR, f"I{top}:I{bot}")
        names[f"cAct1{L}"] = (S_CR, f"I{top}")
        names[f"cAct10{L}"] = (S_CR, f"I{top}:I{top + 9}")
        y.cell(1, "合計", F_KEY, fill=BG_KEY)
        y.ws.merge_cells(start_row=y.r, start_column=1, end_row=y.r, end_column=4)
        for col in "EFGHI":
            y.cell(" EFGHI".index(col) + 4, f"=SUM({col}{top}:{col}{bot})", F_KEY, YEN, BG_KEY)
        y.r += 1
    
    # ───────── 収入が減ったら ─────────
    z = Sheet(wb, S_RISK, names, [36, 16, 16, 16, 40])
    z.title_row("⑤ 収入が減ったら", "年収が下がった場合の毎月の手取りと、住居費・生活費・他の借入の返済を比べます（税額は概算）。")
    for L in "AB":
        z.section(f"物件{L}", ["項目", "今の年収", "2割減", "4割減", "説明"])
        z.ws.cell(row=z.r - 2, column=2).value = z.f(f"{{label{L}}}")
        z.ws.cell(row=z.r - 2, column=2).font = F_LINK
        for key, label, expr, fmt, note in [
            ("rg", "年収", "{g@}", YEN, ""),
            ("rh", "手取り（月額・概算）", "{thm@}", YEN, ""),
            ("rl", "返済÷手取り", "IF({thm@}=0,0,{pm" + L + "}/{thm@})", PCT, ""),
            ("rho", "住居費÷手取り", "IF({thm@}=0,0,{om" + L + "}/{thm@})", PCT, ""),
            ("rgap", "毎月の過不足", "{thm@}-({om" + L + "}+{living}+{existing})", YEN, "手取り−（住居費＋生活費＋他の借入の返済）"),
            ("rmo", "預貯金で補える月数", 'IF({rgap' + L + '@}>=0,"不足なし",MAX(0,{cl' + L + '})/-{rgap' + L + '@})', NUM, "マイナスのとき、手元に残る預貯金で何か月補えるか"),
        ]:
            z.cell(1, label)
            for k in range(3):
                z.name(f"{key}{L}{k}", 2 + k)
                z.cell(2 + k, z.f(expr.replace("@", str(k))), F_LINK if key in ("rg", "rh") else F_CALC, fmt)
            if note:
                z.cell(5, note, F_NOTE)
            z.r += 1

    # ───────── 売却の目安 ─────────
    v = Sheet(wb, S_SALE, names, [30, 16, 16, 16, 44])
    v.title_row("⑥ 売却の目安（シナリオ）",
                "土地は毎年一定の割合で変化、建物は10年後に ② 詳しい条件 の割合になる直線で減価（下限は今の1割）と仮定。将来の価格を約束するものではありません。")
    v.section("シナリオの前提", ["項目", "下落", "横ばい", "上昇", "メモ"])
    v.cell(1, "地価の変化（年）")
    for k, val in enumerate([-0.01, 0, 0.01]):
        v.cell(2 + k, val, F_IN, PCT, BG_IN)
        v.name(f"lr{k}", 2 + k)
    v.r += 1
    v.cell(1, "10年後の建物価値の増減（入力値に対して）")
    for k, val in enumerate([-0.1, 0, 0.1]):
        v.cell(2 + k, val, F_IN, PCT, BG_IN)
        v.name(f"bdl{k}", 2 + k)
    v.r += 1
    v.cell(1, "建物価値の下限（今の価値に対する割合）")
    v.cell(2, 0.1, F_IN, PCT, BG_IN)
    v.name("bFloor", 2)
    v.r += 1
    v.cell(1, "仲介手数料以外の売却費用")
    v.cell(2, 100_000, F_IN, YEN, BG_IN)
    v.name("saleOther", 2)
    v.cell(5, "登記・書類など（仮定）", F_NOTE)
    v.r += 1
    v.cell(1, "仲介手数料の消費税")
    v.cell(2, 0.10, F_IN, PCT, BG_IN)
    v.name("vat", 2)
    v.cell(5, "仲介手数料は宅地建物取引業法の上限で計算", F_NOTE)
    v.r += 1
    for L in "AB":
        for yrs in (5, 10, 20):
            v.section(f"物件{L}・{yrs}年後に売る場合", ["項目", "下落", "横ばい", "上昇", "説明"])
            v.ws.cell(row=v.r - 2, column=2).value = v.f(f"{{label{L}}}")
            v.ws.cell(row=v.r - 2, column=2).font = F_LINK
            base = v.r
            for key, label, expr, note in [
                ("va", "土地の価値", f"{{land{L}}}*(1+{{lr@}})^{yrs}", ""),
                ("vb", "建物の価値", f"{{bldg{L}}}*MAX({{bFloor}},1-(1-MIN(0.95,MAX(0.05,{{b10{L}}}+{{bdl@}})))/10*{yrs})", ""),
                ("vt", "売却価格の目安", "{va@}+{vb@}", ""),
                ("vc", "売却の費用", "IF({vt@}<=2000000,{vt@}*0.05,IF({vt@}<=4000000,{vt@}*0.04+20000,{vt@}*0.03+60000))*(1+{vat})+{saleOther}", "仲介手数料（上限）＋その他"),
                ("vl", "ローン残高", f"IF({{rate}}=0,MAX(0,{{ln{L}}}-{{pm{L}}}*MIN({yrs}*12,{{years}}*12)),MAX(0,{{ln{L}}}*(1+{{rate}}/12)^MIN({yrs}*12,{{years}}*12)-{{pm{L}}}*((1+{{rate}}/12)^MIN({yrs}*12,{{years}}*12)-1)/({{rate}}/12)))", ""),
                ("vn", "手元に残る額（マイナスは不足）", "{vt@}-{vc@}-{vl@}", "売却価格−費用−ローン残高"),
            ]:
                v.cell(1, label, F_KEY if key == "vn" else F_CALC, fill=BG_KEY if key == "vn" else None)
                for k in range(3):
                    v.name(f"{key}{L}{yrs}_{k}", 2 + k)
                    v.cell(2 + k, v.f(re.sub(r"\{(va|vb|vt|vc|vl)(\d)\}", lambda mm: f"{{{mm.group(1)}{L}{yrs}_{mm.group(2)}}}", expr.replace("@", str(k)))), F_KEY if key == "vn" else F_CALC, YEN, BG_KEY if key == "vn" else None)
                if note:
                    v.cell(5, note, F_NOTE)
                v.r += 1
            names[f"sale{L}{yrs}"] = (S_SALE, f"B{v.r - 1}:D{v.r - 1}")

    # タブの順：①〜⑥（見る順）→ 参考
    order = [S_EASY, S_IN, S_CALC, S_CR, S_RISK, S_SALE, S_TAX, S_TAXT, S_RULE]
    wb._sheets = [wb[n] for n in order]
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = TAB[ws.title]
    wb.active = 0

    pat = re.compile(r"\{([^{}]+)\}")
    for ws in wb.worksheets:
        for rw in ws.iter_rows():
            for cl in rw:
                if isinstance(cl.value, str) and cl.value.startswith("=") and "{" in cl.value:
                    cl.value = pat.sub(lambda mm: _ref(names, mm.group(1), ws.title), cl.value)

    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"資金計画シミュレーター_{ed['label']}.xlsx")
    wb.save(path)
    return path, names


def main():
    data = editions()
    ids = sys.argv[1:] or list(data["E"].keys())
    for id_ in ids:
        ed = data["E"][id_]
        path, _ = build(ed, data["R"], os.path.join(ROOT, "dist", id_))
        print(os.path.relpath(path, ROOT))


if __name__ == "__main__":
    main()
